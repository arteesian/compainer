# services/user_service.py
from database.user_repo import UserRepository
from database.schemas import UserResponse, UserFilter, UserRoleUpdate
from typing import List

import io
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

class UserService:


    @staticmethod
    async def list_users(filters: UserFilter) -> List[UserResponse]:
        users = await UserRepository.get_all_users(filters.role)
        return [UserResponse.model_validate(user) for user in users]


    @staticmethod
    async def update_user_role(update_data: UserRoleUpdate) -> bool:
        return await UserRepository.set_role(
            email=update_data.email,
            is_admin=update_data.is_admin,
            is_vip=update_data.is_vip
        )

    @staticmethod
    async def delete_user(email: str) -> bool:
        return await UserRepository.delete_user(email)

    @staticmethod
    async def export_users_to_excel(filters: UserFilter) -> bytes:
        """
        Возвращает содержимое xlsx-файла с пользователями в виде байтов.
        Фильтрация такая же, как в списке (по role).
        """
        users = await UserRepository.get_all_users(filters.role)

        wb = Workbook()
        ws = wb.active
        ws.title = "Users"

        # Шапка таблицы — здесь можно добавить/убрать поля по вкусу
        headers = [
            "Почта сотрудника",
            "VIP?",
            "Админ?",
            "Супер-админ?",
            "Последняя активность",
        ]
        ws.append(headers)

        for u in users:
            last_activity = u.last_activity_at

            # убираем tzinfo, чтобы openpyxl не падал
            if last_activity is not None and last_activity.tzinfo is not None:
                local_tz = ZoneInfo("Europe/Moscow")
                last_activity = last_activity.astimezone(local_tz).replace(tzinfo=None)

            row = [
                u.email,
                "Да" if u.is_vip else "Нет",
                "Да" if u.is_admin else "Нет",
                "Да" if getattr(u, "is_superadmin") else "Нет",
                last_activity,
            ]
            ws.append(row)


            # если в этой строке есть дата – задаём ей читаемый формат
            if last_activity is not None:
                cell = ws.cell(row=ws.max_row, column=5)  # 5-й столбец
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

        for column_cells in ws.columns:
            max_length = 0
            column = column_cells[0].column  # номер колонки (1,2,3...)
            for cell in column_cells:
                value = cell.value
                if value is None:
                    continue

                # для datetime приводим к строке в том же формате
                if isinstance(value, datetime):
                    value_str = value.strftime("%Y-%m-%d %H:%M:%S")
                else:
                    value_str = str(value)

                max_length = max(max_length, len(value_str))

            column_letter = get_column_letter(column)
            # +2, чтобы был небольшой запас
            ws.column_dimensions[column_letter].width = max_length + 2

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream.getvalue()